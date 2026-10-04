#!/usr/bin/env python3
"""Build a publication-facing hash manifest for frozen QSS evidence artifacts."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def resolve(value: str | Path) -> Path:
    path = Path(value).expanduser()
    return path if path.is_absolute() else PROJECT_ROOT / path


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def schema_label(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".json":
            payload = json.loads(path.read_text(encoding="utf-8"))
            version = payload.get("schema_version") if isinstance(payload, dict) else None
            return f"json-schema-{version}" if version is not None else "json-schema-not-declared"
        if suffix == ".jsonl":
            with path.open(encoding="utf-8") as handle:
                first = next((json.loads(line) for line in handle if line.strip()), {})
            keys = "\x1f".join(sorted(first)) if isinstance(first, dict) else ""
            return f"jsonl-keys-{hashlib.sha256(keys.encode()).hexdigest()[:12]}"
        if suffix == ".csv":
            with path.open(encoding="utf-8", newline="") as handle:
                header = next(csv.reader(handle), [])
            keys = "\x1f".join(header)
            return f"csv-header-{hashlib.sha256(keys.encode()).hexdigest()[:12]}"
        if suffix == ".bib":
            return "bibtex"
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                if "Paper Review" not in workbook.sheetnames:
                    return "xlsx-openxml"
                ws = workbook["Paper Review"]
                prefix = [ws.cell(1, column).value for column in range(1, 8)]
                canonical = [
                    "review_row_id", "title", "abstract", "year",
                    "venue", "doi", "openalex_id",
                ]
                return "xlsx-qss-paper-review-v1" if prefix == canonical else "xlsx-openxml"
            finally:
                workbook.close()
        if suffix == ".md":
            return "markdown"
        if suffix == ".py":
            return "python-source"
    except (OSError, ValueError, StopIteration, json.JSONDecodeError):
        return f"{suffix.lstrip('.') or 'binary'}-unreadable-schema"
    return suffix.lstrip(".") or "binary"


def logical_path(path: Path, corpus: str | None = None) -> str:
    try:
        return path.resolve().relative_to(PROJECT_ROOT.resolve()).as_posix()
    except ValueError:
        label = corpus or "shared"
        return f"external_source/{label}/{path.name}"


def release_status(path: Path, role: str) -> tuple[str, str]:
    try:
        path.resolve().relative_to(PROJECT_ROOT.resolve())
        internal = True
    except ValueError:
        internal = False
    if not internal:
        return (
            "private_source_not_for_release",
            "Hash for local verification only; obtain owner/publisher permission before redistribution.",
        )
    if path.name == "REQUEST_MANIFEST.json":
        return (
            "sanitize_paths_before_release",
            "Contains machine-local source paths; publish a redacted copy while retaining this hash privately.",
        )
    if role in {
        "candidate_pool", "retrieval_log", "human_review_queue",
        "human_review_workbook", "self_hosting_candidates",
    }:
        return (
            "metadata_license_review_before_release",
            "Contains scholarly metadata and/or reconstructed abstracts; check redistribution terms and redact if needed.",
        )
    return (
        "public_release_candidate",
        "Generated evidence/configuration suitable for the reproducibility package after final manuscript QA.",
    )


def artifact_row(
    path: Path,
    scope: str,
    role: str,
    corpus: str | None = None,
    cutoff: int | None = None,
) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(path)
    status, note = release_status(path, role)
    return {
        "scope": scope,
        "corpus": corpus or "",
        "cutoff": cutoff if cutoff is not None else "",
        "role": role,
        "path": logical_path(path, corpus),
        "sha256": sha256(path),
        "bytes": path.stat().st_size,
        "schema_or_format_version": schema_label(path),
        "release_status": status,
        "release_note": note,
    }


def add(
    rows: list[dict[str, Any]], seen: set[Path], path: Path, scope: str, role: str,
    corpus: str | None = None, cutoff: int | None = None,
) -> None:
    key = path.resolve()
    if key in seen:
        return
    rows.append(artifact_row(path, scope, role, corpus, cutoff))
    seen.add(key)


def collect(registry_path: Path) -> list[dict[str, Any]]:
    registry = json.loads(registry_path.read_text(encoding="utf-8"))
    rows: list[dict[str, Any]] = []
    seen: set[Path] = set()
    add(rows, seen, registry_path, "core_evaluation", "corpus_registry")
    for path in (
        PROJECT_ROOT / "config/qss_late_pruner_v2_frozen.json",
        PROJECT_ROOT / "scripts/evaluate_qss_late_pruner.py",
        PROJECT_ROOT / "scripts/derive_qss_stage_diagnostics.py",
        PROJECT_ROOT / "docs/qss_reproducibility_supplement_2026-08-28.md",
    ):
        role = "reproducibility_documentation" if path.suffix == ".md" else "frozen_policy_or_runner"
        add(rows, seen, path, "core_evaluation", role)

    for corpus, entry in registry.get("corpora", {}).items():
        validation = entry.get("qss_validation") or {}
        if not validation or corpus == "madlab_template":
            continue
        # Topic-formulation diagnostics are deliberately unfrozen design work.
        # Do not mislabel them as frozen evaluation evidence or silently change
        # the publication-facing inventory before the scope-freeze gate passes.
        if validation.get("evaluation_mode") == "topic_formulation_scope_diagnostic":
            continue
        study = resolve(validation["study_config"])
        add(rows, seen, study, "core_evaluation", "frozen_study_config", corpus)
        bib = resolve(validation.get("evaluation_bib", entry["bib"]))
        add(rows, seen, bib, "core_evaluation", "evaluation_bibliography", corpus)
        for optional_key, role in (
            ("seed_scope_audit", "seed_scope_audit"),
            ("seed_expansion_manifest", "seed_expansion_manifest"),
        ):
            if validation.get(optional_key):
                add(rows, seen, resolve(validation[optional_key]), "core_evaluation", role, corpus)
        for cutoff_value in validation.get("cutoffs", []):
            cutoff = int(cutoff_value)
            candidate_value = validation.get("candidate_pools", {}).get(str(cutoff))
            candidate = (
                resolve(candidate_value) if candidate_value else
                PROJECT_ROOT / "outputs/qss-multicorpus" / corpus / "candidates" / f"qss_candidates_{cutoff}.json"
            )
            evaluation_value = validation.get("evaluation_outputs", {}).get(str(cutoff))
            evaluation = (
                resolve(evaluation_value) if evaluation_value else
                PROJECT_ROOT / "outputs/qss-multicorpus" / corpus / "evaluation" / str(cutoff)
            )
            add(rows, seen, candidate, "core_evaluation", "candidate_pool", corpus, cutoff)
            for name, role in (
                ("REQUEST_MANIFEST.json", "evaluation_request_manifest"),
                ("summary.json", "evaluation_summary"),
                ("per_bag_results.csv", "per_bag_results"),
                ("late_pruning_audit.jsonl", "late_pruning_audit"),
            ):
                add(rows, seen, evaluation / name, "core_evaluation", role, corpus, cutoff)

    diagnostics = PROJECT_ROOT / "outputs/qss-stage-diagnostics"
    for name, role in (
        ("stage_diagnostics.json", "derived_diagnostics"),
        ("case_summary.csv", "derived_case_summary"),
        ("group_stage_localization.csv", "derived_group_localization"),
        ("exact_positive_selection_frequency.csv", "derived_exact_frequency"),
        ("selection_set_stability.csv", "derived_stability"),
        ("discovery_origin_concentration.csv", "derived_origin_concentration"),
        ("human_construct_audit_queue.csv", "human_review_queue"),
        ("REPORT.md", "derived_report"),
        ("HUMAN_AUDIT_GUIDE.md", "audit_protocol"),
    ):
        add(rows, seen, diagnostics / name, "derived_diagnostics", role)

    self_host = PROJECT_ROOT / "outputs/qss-self-hosting/round0"
    add(
        rows, seen, PROJECT_ROOT / "config/qss_self_hosting_concept_frozen.json",
        "self_hosting_supplement", "frozen_concept_package",
    )
    add(
        rows, seen, PROJECT_ROOT / "scripts/run_qss_self_hosting_round0.py",
        "self_hosting_supplement", "self_hosting_runner",
    )
    for name, role in (
        ("REQUEST_MANIFEST.json", "self_hosting_request_manifest"),
        ("SUMMARY.json", "self_hosting_summary"),
        ("round0_candidates.json", "self_hosting_candidates"),
        ("round0_review_queue.csv", "human_review_queue"),
        ("OPENALEX_REQUESTS.jsonl", "retrieval_log"),
        ("REPORT.md", "self_hosting_report"),
        ("REVIEW_GUIDE.md", "self_hosting_review_protocol"),
    ):
        add(rows, seen, self_host / name, "self_hosting_supplement", role)

    workbook_config = PROJECT_ROOT / "config/qss_review_workbook_experiment.json"
    workbook_pipeline = PROJECT_ROOT / "scripts/run_qss_review_workbook_pipeline.py"
    workbook_validation = PROJECT_ROOT / "scripts/validate_qss_review_workbooks.py"
    workbook_output = PROJECT_ROOT / "outputs/qss-review-workbook-experiment"
    for path, role in (
        (workbook_config, "workbook_experiment_config"),
        (workbook_pipeline, "workbook_experiment_runner"),
        (workbook_validation, "workbook_contract_validator"),
        (workbook_output / "validation_report.json", "workbook_validation_report"),
        (workbook_output / "RUN_MANIFEST.json", "workbook_experiment_manifest"),
    ):
        add(rows, seen, path, "human_review_instrumentation", role)
    workbook_spec = json.loads(workbook_config.read_text(encoding="utf-8"))
    for item in workbook_spec["declared_outputs"]:
        add(
            rows, seen, resolve(item["path"]), "human_review_instrumentation",
            "human_review_workbook",
        )
    return sorted(rows, key=lambda row: (row["scope"], row["corpus"], str(row["cutoff"]), row["role"], row["path"]))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def readme(rows: list[dict[str, Any]]) -> str:
    statuses: dict[str, int] = {}
    for row in rows:
        statuses[row["release_status"]] = statuses.get(row["release_status"], 0) + 1
    status_lines = "\n".join(f"- `{key}`: {value}" for key, value in sorted(statuses.items()))
    return f"""# QSS Reproducibility Artifact Manifest

This publication-facing inventory hashes {len(rows)} frozen or derived evidence
artifacts. It was generated offline and does not alter evaluation outcomes.

## Files

- `artifact_manifest.json`: machine-readable manifest, schema version 1.
- `artifact_manifest.csv`: flat review copy with the same artifact rows.

## Release Status Summary

{status_lines}

## Interpretation

- Project-relative paths are safe logical identifiers. External bibliographies
  are represented as `external_source/<corpus>/<basename>` so the public
  manifest does not expose private absolute paths.
- `private_source_not_for_release` artifacts are hashed for local verification
  only and must not be bundled without permission.
- Request manifests need path sanitization before release.
- Candidate/retrieval/human-review metadata needs a redistribution-license and
  abstract-content check before release.
- A SHA-256 change means the artifact no longer matches this frozen inventory;
  rebuild the manifest only as an explicit new evidence version.
- Exploratory topic-formulation cases are intentionally excluded until their
  evaluation object passes a scope-freeze gate.
"""


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--registry", type=Path,
        default=PROJECT_ROOT / "config/qss_corpus_registry.json",
    )
    parser.add_argument(
        "--output-dir", type=Path,
        default=PROJECT_ROOT / "outputs/qss-reproducibility-manifest",
    )
    args = parser.parse_args()
    rows = collect(args.registry)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema_version": 1,
        "hash_algorithm": "sha256",
        "artifact_count": len(rows),
        "artifacts": rows,
    }
    (args.output_dir / "artifact_manifest.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    write_csv(args.output_dir / "artifact_manifest.csv", rows)
    (args.output_dir / "README.md").write_text(readme(rows), encoding="utf-8")
    print(f"Wrote {len(rows)} artifacts to {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
